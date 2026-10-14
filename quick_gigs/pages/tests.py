import io
import json

from django.contrib.auth.models import User
from django.contrib.admin.sites import AdminSite
from django.core.files.uploadedfile import SimpleUploadedFile
from django.contrib.sites.models import Site
from django.test import RequestFactory, SimpleTestCase, TestCase
from allauth.socialaccount.models import SocialApp

from pages.admin import UserProfileAdmin
from pages.models import Application, Gig, SkillLearningProgress, UserProfile

from pages.skillmatch import (
    extract_skills,
    calculate_match_score,
    recommend_jobs,
    build_skill_gap_analysis,
    build_resume_from_experience,
)


class SkillMatchCoreTests(SimpleTestCase):
    def test_extract_skills_handles_english_examples(self):
        skills = extract_skills("I worked in a hotel for two years. I cleaned rooms and helped customers with billing.")
        assert "housekeeping" in skills
        assert "customer service" in skills
        assert "basic billing" in skills

    def test_extract_skills_handles_tamil_examples(self):
        skills = extract_skills("எனக்கு ஹோட்டலில் இரண்டு வருட அனுபவம் உள்ளது. நான் அறைகளை சுத்தம் செய்தேன்.")
        assert "housekeeping" in skills
        assert "hospitality" in skills

    def test_match_score_is_calculated_and_reasonable(self):
        match = calculate_match_score(
            required_skills=["housekeeping", "customer service"],
            user_skills=["housekeeping", "customer service", "communication"],
            preferences={"location": "Bengaluru", "language": "English"},
            job_location="Bengaluru",
            job_language="English",
        )
        assert 70 <= match <= 100

    def test_recommend_jobs_returns_ranked_matches(self):
        jobs = recommend_jobs(
            ["housekeeping", "customer service", "communication"],
            location="Bengaluru",
            preferred_language="English",
        )
        assert jobs
        assert jobs[0]["title"]
        assert jobs[0]["match_score"] >= 60

    def test_skill_gap_analysis_identifies_missing_skills(self):
        job = {
            "title": "Housekeeping Associate",
            "required_skills": ["housekeeping", "customer service", "communication"],
            "preferred_skills": ["basic billing"],
            "description": "Support rooms and guest service.",
        }
        analysis = build_skill_gap_analysis(
            ["housekeeping", "customer service"],
            job,
            preferences={"location": "Bengaluru", "language": "English"},
        )
        assert analysis["match_score"] >= 60
        assert "communication" in analysis["missing_skills"]
        assert analysis["matched_skills"]
        assert analysis["explanation"]

    def test_resume_generation_uses_experience_text_without_runtime_error(self):
        resume = build_resume_from_experience(
            "I worked in a hotel for two years. I cleaned rooms and served customers.",
            {"full_name": "Demo Applicant", "location": "Bengaluru", "language": "English"},
        )
        assert resume["professional_summary"]
        assert "housekeeping" in resume["skills"]
        assert resume["experience"][0]["details"]


class SkillMatchWebFeatureTests(TestCase):
    def setUp(self):
        google_app = SocialApp.objects.create(
            provider="google",
            name="Google test",
            client_id="test-client",
            secret="test-secret",
        )
        google_app.sites.add(Site.objects.get_current())
        self.user = User.objects.create_user(
            username="candidate",
            email="candidate@example.com",
            password="test-password",
        )
        UserProfile.objects.create(
            user=self.user,
            full_name="Candidate Example",
            phone_number="9876543210",
            account_type="worker",
            skills="housekeeping, customer service",
            verification_status="approved",
        )
        self.job = Gig.objects.create(
            employer=User.objects.create_user(username="employer", password="test-password"),
            title="Room Attendant",
            description="Clean guest rooms and support guests.",
            employer_name="Local Hotel",
            location="Chennai",
            pay="₹500/day",
            skills="housekeeping, communication",
            status="open",
        )

    def test_registration_and_login_routes_remain_functional(self):
        response = self.client.post(
            "/signup/",
            {
                "full_name": "New Candidate",
                "email": "new-candidate@example.com",
                "username": "new-candidate",
                "phone": "9876543212",
                "password": "ValidPassword123!",
                "confirm_password": "ValidPassword123!",
                "account_type": "worker",
                "skills": "tailoring",
            },
        )
        self.assertRedirects(response, "/login/", fetch_redirect_response=False)
        registered = User.objects.get(username="new-candidate")
        self.assertEqual(registered.profile.verification_status, "pending")
        self.assertFalse(registered.is_active)
        self.assertTrue(registered.check_password("ValidPassword123!"))
        self.assertNotEqual(registered.password, "ValidPassword123!")

        response = self.client.post(
            "/login/",
            {"email": "new-candidate@example.com", "password": "ValidPassword123!"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "verification-status.html")
        self.assertContains(response, "Verification pending")
        self.assertFalse(response.wsgi_request.user.is_authenticated)

        registered.profile.verification_status = "approved"
        profile_admin = UserProfileAdmin(UserProfile, AdminSite())
        profile_admin.save_model(
            RequestFactory().post("/admin/pages/userprofile/"),
            registered.profile,
            form=None,
            change=True,
        )
        registered.refresh_from_db()
        self.assertTrue(registered.is_active)
        self.assertFalse(registered.profile.aadhaar_verified)
        response = self.client.post(
            "/login/",
            {"email": "new-candidate@example.com", "password": "ValidPassword123!"},
        )
        self.assertRedirects(response, "/", fetch_redirect_response=False)
        self.assertTrue(response.wsgi_request.user.is_authenticated)
        response = self.client.get("/logout/")
        self.assertRedirects(response, "/", fetch_redirect_response=False)
        self.assertRedirects(self.client.get("/profile/"), "/login/?next=/profile/")

    def test_rejected_user_sees_reason_and_pending_users_cannot_use_protected_pages(self):
        rejected = User.objects.create_user(
            username="rejected-candidate",
            email="rejected@example.com",
            password="ValidPassword123!",
            is_active=False,
        )
        UserProfile.objects.create(
            user=rejected,
            full_name="Rejected Candidate",
            phone_number="9876543213",
            account_type="worker",
            verification_status="rejected",
            rejection_reason="Please provide a clearer document image.",
        )

        response = self.client.post(
            "/login/",
            {"email": "rejected@example.com", "password": "ValidPassword123!"},
        )
        self.assertTemplateUsed(response, "verification-status.html")
        self.assertContains(response, "Verification needs attention")
        self.assertContains(response, "Please provide a clearer document image.")
        self.assertFalse(response.wsgi_request.user.is_authenticated)

        pending = User.objects.create_user(
            username="pending-candidate",
            email="pending@example.com",
            password="ValidPassword123!",
            is_active=True,
        )
        UserProfile.objects.create(
            user=pending,
            full_name="Pending Candidate",
            phone_number="9876543214",
            account_type="employer",
            verification_status="pending",
        )
        self.client.force_login(pending)
        self.assertRedirects(self.client.get("/"), "/verification-status/")
        response = self.client.get("/profile/")
        self.assertRedirects(response, "/verification-status/")
        self.assertContains(self.client.get("/verification-status/"), "Verification pending")

    def test_login_rejects_bad_password_and_uses_saved_account_type_for_redirect(self):
        response = self.client.post(
            "/login/",
            {"email": "candidate@example.com", "password": "incorrect"},
        )
        self.assertContains(response, "Invalid email or password.")

        employer = User.objects.create_user(
            username="approved-employer",
            email="employer@example.com",
            password="EmployerPassword123!",
        )
        UserProfile.objects.create(
            user=employer,
            full_name="Approved Employer",
            phone_number="9876543215",
            account_type="employer",
            verification_status="approved",
        )
        response = self.client.post(
            "/login/",
            {"email": "employer@example.com", "password": "EmployerPassword123!"},
        )
        self.assertRedirects(response, "/", fetch_redirect_response=False)
        self.assertTrue(response.wsgi_request.user.is_authenticated)
        self.assertEqual(response.wsgi_request.user.profile.account_type, "employer")

    def test_signup_rejects_invalid_account_type_instead_of_defaulting_to_worker(self):
        response = self.client.post(
            "/signup/",
            {
                "full_name": "Unknown Role",
                "email": "unknown-role@example.com",
                "username": "unknown-role",
                "phone": "9876543216",
                "password": "ValidPassword123!",
                "confirm_password": "ValidPassword123!",
                "account_type": "invalid",
            },
        )
        self.assertContains(response, "Please select a valid account type.")
        self.assertFalse(User.objects.filter(username="unknown-role").exists())

    def test_navigation_routes_and_password_reset_link_resolve(self):
        response = self.client.get("/login/")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "/accounts/password/reset/")
        for path in ("/", "/signup/", "/job-matches/", "/about/", "/faq/", "/help/", "/safety/", "/terms/", "/contact/"):
            with self.subTest(path=path):
                response = self.client.get(path)
                if path == "/faq/":
                    self.assertRedirects(response, "/help-center/")
                else:
                    self.assertEqual(response.status_code, 200)

        help_page = self.client.get("/help/")
        self.assertContains(help_page, "/safety-guidelines/")
        self.assertContains(help_page, "/contact-support/")


    def test_job_match_page_uses_real_open_gigs_and_profile_skills(self):
        self.client.force_login(self.user)
        response = self.client.get("/job-matches/")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Room Attendant")
        self.assertContains(response, "50%")
        self.assertContains(response, "communication")
        self.assertContains(response, "housekeeping, customer service")

    def test_job_matching_and_listings_render_every_available_database_job(self):
        additional = Gig.objects.create(
            employer=self.job.employer,
            title="Housekeeping Supervisor",
            description="Coordinate housekeeping tasks.",
            employer_name="Local Hotel",
            location="Bengaluru",
            pay="₹800/day",
            skills="housekeeping, inventory management",
            status="urgent",
        )
        Gig.objects.create(
            employer=self.job.employer,
            title="Closed historical job",
            description="No longer available.",
            employer_name="Local Hotel",
            location="Chennai",
            pay="₹500/day",
            skills="housekeeping",
            status="closed",
        )
        self.client.force_login(self.user)
        matches = self.client.get("/job-matches/")
        self.assertContains(matches, "Room Attendant")
        self.assertContains(matches, "Housekeeping Supervisor")
        self.assertNotContains(matches, "Closed historical job")
        listings = self.client.get("/jobs/")
        self.assertContains(listings, "Room Attendant")
        self.assertContains(listings, "Housekeeping Supervisor")
        self.assertContains(listings, "Closed historical job", count=0)

    def test_job_detail_shows_authenticated_users_skill_comparison(self):
        self.client.force_login(self.user)
        response = self.client.get(f"/jobs/{self.job.id}/")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Your skill match")
        self.assertContains(response, "50% match")
        self.assertContains(response, "Skills you have:")
        self.assertContains(response, "Skills to learn:")
        self.assertContains(response, "Learn missing skills")

    def test_unmatched_skill_profile_recalculates_job_match(self):
        self.client.force_login(self.user)
        response = self.client.get(
            "/job-matches/",
            {"skills": "housekeeping, communication"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "100%")

    def test_match_scoring_normalizes_skill_punctuation_and_spacing(self):
        self.job.skills = "Customer-Service"
        self.job.save(update_fields=["skills"])
        self.client.force_login(self.user)
        response = self.client.get("/job-matches/", {"skills": "customer service"})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "100%")

    def test_zero_skill_match_is_reported_as_zero(self):
        self.client.force_login(self.user)
        response = self.client.get("/job-matches/", {"skills": "tailoring"})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "0%")
        self.assertContains(response, "housekeeping")

    def test_job_without_requirements_has_no_percentage(self):
        self.job.skills = ""
        self.job.save(update_fields=["skills"])
        self.client.force_login(self.user)
        response = self.client.get("/job-matches/")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "No listed skill requirements")

    def test_match_api_uses_published_gig_records_and_candidate_skills(self):
        response = self.client.get(
            "/api/jobs/recommendations/",
            {"skills": "housekeeping"},
        )
        self.assertEqual(response.status_code, 200)
        job = response.json()["jobs"][0]
        self.assertEqual(job["id"], self.job.id)
        self.assertEqual(job["match_score"], 50)
        self.assertEqual(job["matched_skills"], ["housekeeping"])
        self.assertEqual(job["missing_skills"], ["communication"])

    def test_match_api_defaults_to_saved_profile_skills(self):
        self.client.force_login(self.user)
        response = self.client.get("/api/jobs/recommendations/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["jobs"][0]["match_score"], 50)

    def test_match_api_ranks_multiple_real_jobs_and_keeps_unscorable_jobs_unscored(self):
        best_job = Gig.objects.create(
            employer=self.job.employer,
            title="Housekeeping Assistant",
            description="Support room cleaning.",
            employer_name="City Hotel",
            location="Chennai",
            pay="₹600/day",
            job_type="part-time",
            skills="housekeeping, customer service",
            status="open",
        )
        no_requirements_job = Gig.objects.create(
            employer=self.job.employer,
            title="General Helper",
            description="Requirements not supplied.",
            employer_name="City Hotel",
            location="Chennai",
            pay="₹500/day",
            skills="",
            status="urgent",
        )
        self.client.force_login(self.user)
        response = self.client.get("/api/jobs/recommendations/")
        self.assertEqual(response.status_code, 200)
        jobs = response.json()["jobs"]
        self.assertEqual(len(jobs), 3)
        self.assertEqual(jobs[0]["id"], best_job.id)
        self.assertEqual(jobs[0]["match_score"], 100)
        self.assertEqual(jobs[0]["company"], "City Hotel")
        self.assertEqual(jobs[0]["job_type"], "part-time")
        self.assertIsNone(jobs[-1]["match_score"])
        self.assertEqual(jobs[-1]["id"], no_requirements_job.id)

    def test_skill_stimulator_selects_target_job_and_restores_learning_progress(self):
        SkillLearningProgress.objects.create(
            user=self.user,
            skill_key="communication",
            skill_name="communication",
            completed=True,
        )
        self.client.force_login(self.user)
        response = self.client.get("/dashboard/", {"job_id": self.job.id})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="stimulator-job"')
        self.assertContains(response, 'data-stimulator-job')
        self.assertContains(response, "Skills to develop")
        self.assertContains(response, "communication")
        self.assertContains(response, "stimulator-completed-skills")
        self.assertContains(response, "Learn missing skills")

    def test_skill_gap_analysis_includes_saved_profile_skills(self):
        self.client.force_login(self.user)
        response = self.client.post(
            "/api/experience/analyze/",
            data=json.dumps({"text": "I work with inventory."}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("housekeeping", response.json()["skills"])
        self.assertIn("customer service", response.json()["skills"])
    def test_profile_skill_save_requires_authentication(self):
        response = self.client.post(
            "/api/profile/skills/",
            data=json.dumps({"skills": ["housekeeping"]}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 302)

    def test_authenticated_profile_skill_save_updates_candidate_profile(self):
        self.client.force_login(self.user)
        response = self.client.post(
            "/api/profile/skills/",
            data=json.dumps({"skills": ["housekeeping", "communication"]}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        self.user.profile.refresh_from_db()
        self.assertEqual(
            self.user.profile.get_skills_list(),
            ["housekeeping", "customer service", "communication"],
        )
        repeated = self.client.post(
            "/api/profile/skills/",
            data=json.dumps({"skills": ["Customer-Service", "communication"]}),
            content_type="application/json",
        )
        self.assertEqual(repeated.status_code, 200)
        self.assertEqual(len(repeated.json()["skills"]), 3)

    def test_learning_progress_is_private_persistent_and_upserted(self):
        payload = {"skill": "Customer-Service", "completed": True}
        response = self.client.post(
            "/api/learning/progress/",
            data=json.dumps(payload),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 302)

        self.client.force_login(self.user)
        response = self.client.post(
            "/api/learning/progress/",
            data=json.dumps(payload),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["completed"])

        response = self.client.post(
            "/api/learning/progress/",
            data=json.dumps({"skill": "customer service", "completed": False}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.json()["completed"])
        self.assertEqual(SkillLearningProgress.objects.filter(user=self.user).count(), 1)

    def test_job_match_restores_saved_learning_progress_after_page_reload(self):
        SkillLearningProgress.objects.create(
            user=self.user,
            skill_key="communication",
            skill_name="communication",
            completed=True,
        )
        self.client.force_login(self.user)
        response = self.client.get("/job-matches/")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="completed-skill-keys-data"')
        self.assertContains(response, '"communication"')

    def test_skill_simulator_projects_real_job_score(self):
        response = self.client.post(
            "/api/matching/simulate/",
            data=json.dumps({
                "user_skills": ["housekeeping"],
                "selected_skill": "communication",
            }),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        simulation = response.json()["simulation"]
        self.assertEqual(simulation["before"], 50)
        self.assertEqual(simulation["after"], 100)
        self.assertEqual(simulation["improvement"], 50)
        self.assertEqual(simulation["jobs_count"], 1)

    def test_skill_simulator_defaults_to_profile_skills(self):
        self.client.force_login(self.user)
        response = self.client.post(
            "/api/matching/simulate/",
            data=json.dumps({"selected_skill": "communication"}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        simulation = response.json()["simulation"]
        self.assertEqual(simulation["before"], 50)
        self.assertEqual(simulation["after"], 100)

    def test_skill_simulator_rejects_a_skill_that_is_not_missing(self):
        response = self.client.post(
            "/api/matching/simulate/",
            data=json.dumps({
                "user_skills": ["housekeeping", "customer service"],
                "selected_skill": "customer service",
            }),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("missing from a published job", response.json()["error"])

    def test_skill_simulation_does_not_claim_a_score_without_job_requirements(self):
        self.job.skills = ""
        self.job.save(update_fields=["skills"])
        response = self.client.post(
            "/api/matching/simulate/",
            data=json.dumps({
                "user_skills": [],
                "selected_skill": "communication",
            }),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        simulation = response.json()["simulation"]
        self.assertIsNone(simulation["before"])
        self.assertIsNone(simulation["after"])
        self.assertEqual(simulation["jobs_count"], 0)

    def test_complete_profile_saves_and_returns_employer_to_original_destination(self):
        self.client.force_login(self.user)
        response = self.client.post(
            "/complete-profile/",
            {
                "full_name": "Updated Candidate",
                "phone_number": "9876543210",
                "account_type": "employer",
                "next": "/post-job/",
            },
        )
        self.assertRedirects(response, "/post-job/", fetch_redirect_response=False)
        self.user.profile.refresh_from_db()
        self.assertEqual(self.user.profile.full_name, "Updated Candidate")
        self.assertEqual(self.user.profile.phone_number, "9876543210")
        self.assertEqual(self.user.profile.account_type, "employer")
        self.assertEqual(self.client.get("/post-job/").status_code, 200)

    def test_complete_profile_worker_does_not_continue_to_post_job(self):
        self.client.force_login(self.user)
        response = self.client.post(
            "/complete-profile/",
            {
                "full_name": "Candidate Example",
                "phone_number": "9876543210",
                "account_type": "worker",
                "next": "/post-job/",
            },
        )
        self.assertRedirects(response, "/dashboard/", fetch_redirect_response=False)
        self.assertEqual(Gig.objects.filter(employer=self.user).count(), 0)

    def test_complete_profile_shows_validation_errors_without_saving(self):
        self.client.force_login(self.user)
        response = self.client.post(
            "/complete-profile/",
            {
                "full_name": "",
                "phone_number": "123",
                "account_type": "invalid",
                "next": "https://example.invalid/",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Full name is required.")
        self.assertContains(response, "valid phone number")
        self.user.profile.refresh_from_db()
        self.assertEqual(self.user.profile.account_type, "worker")

    def test_complete_profile_rejects_external_redirects(self):
        self.client.force_login(self.user)
        response = self.client.post(
            "/complete-profile/",
            {
                "full_name": "Candidate Example",
                "phone_number": "9876543210",
                "account_type": "employer",
                "next": "https://example.invalid/",
            },
        )
        self.assertRedirects(response, "/dashboard/", fetch_redirect_response=False)

    def test_worker_cannot_post_a_job_and_original_destination_is_preserved(self):
        self.client.force_login(self.user)
        response = self.client.post(
            "/post-job/",
            {
                "title": "Unauthorized job",
                "employer_name": "Candidate",
                "location": "Chennai",
                "pay": "₹500/day",
                "description": "Should not be saved.",
            },
        )
        self.assertRedirects(
            response,
            "/complete-profile/?next=%2Fpost-job%2F",
            fetch_redirect_response=False,
        )
        self.assertFalse(Gig.objects.filter(title="Unauthorized job").exists())

    def test_employer_post_saves_supported_fields_and_job_is_listed(self):
        self.client.force_login(self.user)
        self.user.profile.account_type = "employer"
        self.user.profile.save(update_fields=["account_type"])
        response = self.client.post(
            "/post-job/",
            {
                "title": "Hotel Room Attendant",
                "employer_name": "Local Hotel",
                "location": "Chennai",
                "pay": "₹500/day",
                "duration": "Ongoing",
                "job_type": "full-time",
                "schedule": "9 AM to 5 PM",
                "contact_info": "employer@example.com",
                "skills": "housekeeping, communication",
                "description": "Clean guest rooms and support guests.",
            },
        )
        self.assertRedirects(response, "/jobs/", fetch_redirect_response=False)
        posted = Gig.objects.get(employer=self.user, title="Hotel Room Attendant")
        self.assertEqual(posted.employer_name, "Local Hotel")
        self.assertEqual(posted.location, "Chennai")
        self.assertEqual(posted.pay, "₹500/day")
        self.assertEqual(posted.duration, "Ongoing")
        self.assertEqual(posted.job_type, "full-time")
        self.assertEqual(posted.work_schedule, "9 AM to 5 PM")
        self.assertEqual(posted.contact_info, "employer@example.com")
        self.assertEqual(posted.skills, "housekeeping, communication")
        self.assertEqual(posted.status, "open")
        owner_jobs = self.client.get("/jobs/")
        self.assertContains(owner_jobs, "Hotel Room Attendant")
        detail = self.client.get(f"/jobs/{posted.id}/")
        self.assertContains(detail, "Clean guest rooms and support guests.")
        self.assertContains(detail, "Full Time")
        self.assertContains(detail, "employer@example.com")

        worker = User.objects.create_user(username="second-candidate", password="example-password")
        UserProfile.objects.create(
            user=worker,
            full_name="Second Candidate",
            phone_number="9876543211",
            account_type="worker",
            skills="housekeeping",
            verification_status="approved",
        )
        self.client.force_login(worker)
        jobs_response = self.client.get("/jobs/")
        self.assertContains(jobs_response, "Hotel Room Attendant")
        matches_response = self.client.get("/api/jobs/recommendations/", {"skills": "housekeeping"})
        match = next(job for job in matches_response.json()["jobs"] if job["id"] == posted.id)
        self.assertEqual(match["match_score"], 50)

    def test_employer_post_form_displays_validation_errors(self):
        self.client.force_login(self.user)
        self.user.profile.account_type = "employer"
        self.user.profile.save(update_fields=["account_type"])
        response = self.client.post("/post-job/", {})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Job title is required.")
        self.assertContains(response, "Required skills are required.")
        self.assertContains(response, "Employment type is required.")
        self.assertFalse(Gig.objects.filter(employer=self.user).exists())

    def test_resume_generation_uses_only_submitted_resume_information(self):
        response = self.client.post(
            "/api/resumes/generate/",
            data=json.dumps({
                "full_name": "Candidate Example",
                "email": "candidate@example.com",
                "objective": "Seeking a room attendant position.",
                "skills": ["housekeeping"],
                "education": "High school",
                "projects": "Community garden",
                "experience_text": "I cleaned guest rooms.",
                "experience": "Room attendant, 2023-2025",
                "certifications": "First aid",
                "languages": "Tamil, English",
                "target_job": "Room Attendant",
            }),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        resume = response.json()["resume"]
        self.assertEqual(resume["professional_summary"], "Seeking a room attendant position.")
        self.assertEqual(resume["skills"], ["housekeeping"])
        self.assertEqual(resume["education"], "High school")
        self.assertEqual(resume["experience"], "Room attendant, 2023-2025")
        self.assertEqual(resume["certifications"], "First aid")

    def test_resume_pdf_endpoint_returns_a_downloadable_pdf(self):
        response = self.client.post(
            "/api/resumes/pdf/",
            data=json.dumps({
                "full_name": "Candidate Example",
                "objective": "Seeking work.",
                "skills": ["housekeeping"],
            }),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_resume_pdf_rejects_unsupported_unicode_instead_of_corrupting_text(self):
        response = self.client.post(
            "/api/resumes/pdf/",
            data=json.dumps({
                "full_name": "\u0ba4\u0bae\u0bbf\u0bb4\u0bcd",
                "skills": ["housekeeping"],
            }),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("browser's Print / Save as PDF", response.json()["error"])

    def test_resume_upload_extracts_known_skills_from_docx(self):
        from docx import Document

        document = Document()
        document.add_paragraph("Experience")
        document.add_paragraph("I cleaned rooms and helped customers.")
        buffer = io.BytesIO()
        document.save(buffer)
        upload = SimpleUploadedFile(
            "resume.docx",
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        response = self.client.post("/api/resumes/extract/", {"resume": upload})
        self.assertEqual(response.status_code, 200)
        self.assertIn("housekeeping", response.json()["skills"])
        self.assertIn("customer service", response.json()["skills"])
        self.assertIn("I cleaned rooms", response.json()["sections"]["experience"])

    def test_resume_upload_extracts_text_from_pdf(self):
        from reportlab.pdfgen import canvas

        buffer = io.BytesIO()
        pdf = canvas.Canvas(buffer)
        pdf.drawString(50, 750, "Experience")
        pdf.drawString(50, 730, "I cleaned rooms and helped customers.")
        pdf.save()
        upload = SimpleUploadedFile(
            "resume.pdf",
            buffer.getvalue(),
            content_type="application/pdf",
        )
        response = self.client.post("/api/resumes/extract/", {"resume": upload})
        self.assertEqual(response.status_code, 200)
        self.assertIn("housekeeping", response.json()["skills"])
        self.assertIn("customer service", response.json()["skills"])

    def test_resume_upload_rejects_unsupported_files(self):
        upload = SimpleUploadedFile("resume.exe", b"not a resume")
        response = self.client.post("/api/resumes/extract/", {"resume": upload})
        self.assertEqual(response.status_code, 400)

    def test_workforce_dashboard_uses_recorded_activity_without_claiming_hires(self):
        statuses = ("pending", "accepted", "rejected", "completed")
        for index, status in enumerate(statuses):
            if index == 0:
                applicant = self.user
            else:
                applicant = User.objects.create_user(
                    username=f"analytics-worker-{index}",
                    email=f"analytics-worker-{index}@example.com",
                    password="ValidPassword123!",
                )
                UserProfile.objects.create(
                    user=applicant,
                    full_name=f"Analytics Worker {index}",
                    phone_number=f"987654320{index}",
                    account_type="worker",
                    skills="housekeeping, communication",
                    verification_status="approved",
                )
            Application.objects.create(
                gig=self.job,
                applicant=applicant,
                status=status,
            )

        SkillLearningProgress.objects.create(
            user=self.user,
            skill_key="communication",
            skill_name="Communication",
            completed=True,
        )
        SkillLearningProgress.objects.create(
            user=self.user,
            skill_key="housekeeping",
            skill_name="Housekeeping",
            completed=False,
        )

        response = self.client.get("/sdg1-impact/?period=6")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["metrics"]["registered_job_seekers"], 4)
        self.assertEqual(response.context["metrics"]["active_job_listings"], 1)
        self.assertEqual(response.context["metrics"]["job_applications"], 4)
        self.assertEqual(response.context["metrics"]["pending_applications"], 1)
        self.assertEqual(response.context["metrics"]["accepted_applications"], 1)
        self.assertEqual(response.context["metrics"]["rejected_applications"], 1)
        self.assertEqual(response.context["metrics"]["completed_applications"], 1)
        self.assertEqual(response.context["metrics"]["learning_resources_completed"], 1)
        self.assertEqual(response.context["metrics"]["learning_completion_rate"], 50)
        self.assertEqual(response.context["application_job_types"][0]["label"], "Not specified")
        self.assertIsNone(response.context["metrics"]["seeking_status"])
        self.assertIsNone(response.context["metrics"]["employment_status"])
        self.assertIsNone(response.context["metrics"]["verified_hires"])
        self.assertEqual(len(response.context["trends"]), 6)
        self.assertContains(response, "Verified hires through SkillMatch AI")
        self.assertContains(response, "Not tracked")
        self.assertContains(response, "Accepted (not verified hires)")
        self.assertContains(response, "Completed application records (not verified hires)")
        self.assertContains(response, "Job type distribution")

    def test_workforce_dashboard_rejects_unsupported_trend_period(self):
        response = self.client.get("/sdg1-impact/?period=999")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["period"], 12)
        self.assertEqual(len(response.context["trends"]), 12)
